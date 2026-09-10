"""Processa PROGRAMADO e KM Não Realizada com baixo uso de memória."""
from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

COL_EMPRESA = 3
COL_LINHA = 4
COL_FROTA = {"U": 11, "S": 12, "D": 13}
COL_VIAGENS = {"U": 14, "S": 15, "D": 16}
COL_OPERACIONAL = 17
COL_MORTA = 18
COL_TRANSPORTA = 22
CALENDARIO_INICIO = 31
CALENDARIO_FIM = 61
FIELDS = ("frota", "viagens", "km_operacional", "km_morta", "km_transporta", "km_total")


def number(value):
    if value is None or isinstance(value, bool):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def line_key(value):
    """Padroniza 1/001 e preserva sufixos como 001C e 005.6."""
    text = str(value or "").strip().upper()
    if not text:
        return ""
    match = re.fullmatch(r"(\d+)([A-Z].*)?", text)
    return match.group(1).zfill(3) + (match.group(2) or "") if match else text


def is_transporta(line):
    return "transporta" in str(line).casefold()


def dates_in_period(ws, start_date, end_date):
    if not start_date or not end_date or start_date > end_date:
        raise ValueError("Informe um período válido.")
    header = next(ws.iter_rows(min_row=2, max_row=2, max_col=CALENDARIO_FIM, values_only=True), ())
    dates = []
    for column in range(CALENDARIO_INICIO, CALENDARIO_FIM + 1):
        current = to_date(header[column - 1] if len(header) >= column else None)
        if current and start_date <= current <= end_date:
            dates.append((column, current))
    if not dates:
        raise ValueError("Não foram encontradas datas no calendário AE:BI para o período informado.")
    return dates


def calculate_programmed(ws, start_date, end_date):
    dates = dates_in_period(ws, start_date, end_date)
    result = defaultdict(lambda: defaultdict(lambda: {field: 0.0 for field in FIELDS}))
    types = defaultdict(lambda: defaultdict(set))
    transporta_by_company = defaultdict(float)
    planned_by_line = defaultdict(float)

    for row in ws.iter_rows(min_row=3, max_col=CALENDARIO_FIM, values_only=True):
        company = str(row[COL_EMPRESA - 1] if len(row) >= COL_EMPRESA else "").strip()
        if not company or company.casefold() in {"none", "nan", "null", "-", "total", "total geral", "total por empresa"}:
            continue
        line = str(row[COL_LINHA - 1] if len(row) >= COL_LINHA else "").strip()
        transporta = is_transporta(line)
        operational = number(row[COL_OPERACIONAL - 1] if len(row) >= COL_OPERACIONAL else 0)
        dead_rate = number(row[COL_MORTA - 1] if len(row) >= COL_MORTA else 0)
        if transporta:
            transporta_by_company[company] += number(row[COL_TRANSPORTA - 1] if len(row) >= COL_TRANSPORTA else 0)

        for column, current in dates:
            schedule = str(row[column - 1] if len(row) >= column else "").strip().upper()
            if schedule not in COL_FROTA:
                continue
            fleet = number(row[COL_FROTA[schedule] - 1] if len(row) >= COL_FROTA[schedule] else 0)
            trips = number(row[COL_VIAGENS[schedule] - 1] if len(row) >= COL_VIAGENS[schedule] else 0)
            km_operational = 0 if transporta else trips * operational
            km_dead = fleet * dead_rate
            entry = result[company][current.isoformat()]
            entry["frota"] += fleet
            entry["viagens"] += trips
            entry["km_operacional"] += km_operational
            entry["km_morta"] += km_dead
            entry["km_total"] += km_operational + km_dead
            types[company][current.isoformat()].add(schedule)
            if not transporta and line_key(line):
                planned_by_line[(company, line_key(line))] += km_operational

    if not result:
        raise ValueError("Nenhum dado de empresa foi encontrado na aba PROGRAMADO.")
    companies, all_daily, company_daily = [], defaultdict(lambda: {field: 0.0 for field in FIELDS}), {}
    for company, rows in sorted(result.items()):
        total = {"empresa": company, **{field: 0.0 for field in FIELDS}}
        out = []
        for key, values in sorted(rows.items()):
            out.append({"data": key, **values, "tipos": sorted(types[company][key])})
            for field in FIELDS:
                total[field] += values[field]
                all_daily[key][field] += values[field]
        total["km_transporta"] = transporta_by_company[company]
        total["km_total"] += total["km_transporta"]
        companies.append(total)
        company_daily[company] = out
    daily = [{"data": key, **values} for key, values in sorted(all_daily.items())]
    return companies, daily, company_daily, planned_by_line


def find_non_operated_sheet(workbook):
    required = {"data", "empresa", "linha", "km nominal"}
    for ws in workbook.worksheets:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        names = {str(value or "").strip().casefold() for value in header}
        if required.issubset(names):
            return ws
    raise ValueError("A planilha de KM Não Realizada precisa ter as colunas Data, Empresa, Linha e KM NOMINAL.")


def read_non_operated_km(source, start_date, end_date):
    """Lê somente a aba detalhada e somente as quatro colunas necessárias."""
    try:
        workbook = load_workbook(source, data_only=True, read_only=True, keep_links=False)
    except Exception as error:
        raise ValueError("A planilha de KM Não Realizada não é um arquivo Excel válido.") from error
    try:
        ws = find_non_operated_sheet(workbook)
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = {str(value or "").strip().casefold(): index for index, value in enumerate(header)}
        max_col = max(columns[name] for name in ("data", "empresa", "linha", "km nominal")) + 1
        by_line, count = defaultdict(float), 0
        for row in ws.iter_rows(min_row=2, max_col=max_col, values_only=True):
            current = to_date(row[columns["data"]] if len(row) > columns["data"] else None)
            if not current or not start_date <= current <= end_date:
                continue
            company = str(row[columns["empresa"]] if len(row) > columns["empresa"] else "").strip()
            line = line_key(row[columns["linha"]] if len(row) > columns["linha"] else "")
            if not company or not line:
                continue
            by_line[(company, line)] += number(row[columns["km nominal"]] if len(row) > columns["km nominal"] else 0)
            count += 1
        if not count:
            raise ValueError("Não foram encontrados registros de KM Não Realizada no período escolhido.")
        return by_line
    finally:
        workbook.close()


def line_efficiency(planned_by_line, non_operated_by_line):
    rows = []
    for company, line in sorted(set(planned_by_line) | set(non_operated_by_line), key=lambda item: (item[0].casefold(), item[1])):
        planned = planned_by_line.get((company, line), 0.0)
        non_operated = non_operated_by_line.get((company, line), 0.0)
        efficiency = (planned - non_operated) / planned * 100 if planned else None
        rows.append({"empresa": company, "linha": line, "km_programado": planned, "km_nao_realizada": non_operated, "eficiencia": efficiency})
    return sorted(rows, key=lambda item: (item["eficiencia"] is None, item["eficiencia"] if item["eficiencia"] is not None else 999, item["empresa"], item["linha"]))


def line_items(values):
    return [{"empresa": company, "linha": line, "km": km} for (company, line), km in sorted(values.items())]


def line_values(items, field="km"):
    values = defaultdict(float)
    for item in items or []:
        company, line = str(item.get("empresa", "")).strip(), line_key(item.get("linha"))
        if company and line:
            values[(company, line)] += number(item.get(field, item.get("km", 0)))
    return values


def apply_non_operated(payload, non_operated_source, start_date, end_date):
    """Acrescenta ou substitui somente a leitura de KM Não Realizada."""
    planned = line_values(payload.get("line_programmed"))
    if not planned:  # compatibilidade com históricos gerados antes desta melhoria
        planned = line_values(payload.get("lines"), "km_programado")
    if not planned:
        raise ValueError("Este mês não possui uma leitura PROGRAMADO compatível. Atualize primeiro a planilha PROGRAMADO.")
    non_operated = read_non_operated_km(non_operated_source, start_date, end_date)
    payload["line_programmed"] = line_items(planned)
    payload["line_non_operated"] = line_items(non_operated)
    payload["lines"] = line_efficiency(planned, non_operated)
    payload["has_non_operated_data"] = True
    return payload


def safe_name(name, used):
    base = re.sub(r'[\\/:*?\[\]]', " ", name).strip()[:31] or "SEM EMPRESA"
    candidate, index = base, 2
    while candidate.casefold() in used:
        suffix = f" ({index})"
        candidate = base[:31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.casefold())
    return candidate


def write_daily_sheet(book, name, rows, monthly_transporta=None):
    ws = book.create_sheet(name)
    ws.append(["Data", "Frota", "Viagens", "KM Operacional", "KM Morta", "KM Total"])
    for item in rows:
        ws.append([item["data"], item["frota"], item["viagens"], item["km_operacional"], item["km_morta"], item["km_total"]])
    if monthly_transporta is not None:
        totals = {field: sum(number(item.get(field)) for item in rows) for field in ("frota", "viagens", "km_operacional", "km_morta", "km_total")}
        ws.append([])
        ws.append(["KM TRANSPORTA DO MÊS", None, None, None, None, monthly_transporta])
        ws.append(["TOTAL DO MÊS", totals["frota"], totals["viagens"], totals["km_operacional"], totals["km_morta"], totals["km_total"] + monthly_transporta])
        for cell in ws[ws.max_row - 1] + ws[ws.max_row]:
            cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "dd/mm/yyyy"
    for column in "ABCDEF":
        ws.column_dimensions[column].width = 22 if column == "A" else 18


def write_efficiency_sheet(book, rows):
    ws = book.create_sheet("EFICIÊNCIA POR LINHA")
    ws.append(["Empresa", "Linha", "KM Operacional Programado", "KM Não Realizada", "Eficiência"])
    for item in rows:
        ws.append([item["empresa"], item["linha"], item["km_programado"], item["km_nao_realizada"], None if item["eficiencia"] is None else item["eficiencia"] / 100])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "#,##0.00"
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "0.0%"
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for column, width in {"A": 20, "B": 14, "C": 28, "D": 21, "E": 14}.items():
        ws.column_dimensions[column].width = width


def write_report_from_payload(payload, destination: Path):
    output = Workbook()
    output.remove(output.active)
    summary = output.create_sheet("TOTAL POR EMPRESA")
    summary.append(["Empresa", "Frota", "Viagens", "KM Operacional", "KM Morta", "KM Total"])
    for item in payload["companies"]:
        summary.append([item["empresa"], item["frota"], item["viagens"], item["km_operacional"], item["km_morta"], item["km_total"]])
    write_daily_sheet(output, "TOTAL POR DIA", payload["daily"])
    write_efficiency_sheet(output, payload.get("lines", []))
    used = {"total por empresa", "total por dia", "eficiência por linha"}
    company_totals = {item["empresa"]: item for item in payload["companies"]}
    for company, rows in payload.get("company_daily", {}).items():
        write_daily_sheet(output, safe_name(company, used), rows, company_totals[company]["km_transporta"])
    destination.parent.mkdir(parents=True, exist_ok=True)
    output.save(destination)


def build_report_from_file(source, destination: Path, start_date, end_date, non_operated_source=None):
    try:
        workbook = load_workbook(source, data_only=True, read_only=True, keep_links=False)
    except Exception as error:
        raise ValueError("O arquivo PROGRAMADO não é uma planilha Excel válida.") from error
    try:
        if "PROGRAMADO" not in workbook.sheetnames:
            raise ValueError("A planilha precisa ter a aba PROGRAMADO.")
        companies, daily, company_daily, planned_by_line = calculate_programmed(workbook["PROGRAMADO"], start_date, end_date)
    finally:
        workbook.close()
    payload = {"companies": companies, "daily": daily, "company_daily": company_daily, "line_programmed": line_items(planned_by_line), "line_non_operated": [], "lines": line_efficiency(planned_by_line, {}), "has_non_operated_data": False, "parameters": {"bm1": start_date.isoformat(), "bn1": end_date.isoformat()}}
    if non_operated_source:
        apply_non_operated(payload, non_operated_source, start_date, end_date)
    write_report_from_payload(payload, destination)
    return payload


def download_source(sheet_id):
    response = requests.get(f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx", timeout=90)
    response.raise_for_status()
    return response.content


def build_report_from_bytes(source, destination: Path, start_date, end_date):
    return build_report_from_file(io.BytesIO(source), destination, start_date, end_date)


def build_report(sheet_id, destination: Path, start_date, end_date):
    return build_report_from_bytes(download_source(sheet_id), destination, start_date, end_date)
